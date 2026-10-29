try:
    import RPi.GPIO as GPIO
except ImportError:
    GPIO = None

import time
import threading
import sqlite3
from io import BytesIO
from contextlib import closing
from datetime import datetime
from werkzeug.utils import secure_filename
from flask import Flask, render_template, request, redirect, send_file
import json
import os
import uuid
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

app = Flask(__name__)
BLINK_PIN = 18

if GPIO is not None:
    GPIO.setmode(GPIO.BCM)
    GPIO.setup(
        BLINK_PIN,
        GPIO.OUT,
        initial=GPIO.HIGH
    )
else:
    print("RPi.GPIO no disponible; GPIO en modo simulado.")


# ----------------------------------
# Helper Functions
# ----------------------------------

def recipe_file(process_name):

    filename = process_name.replace(" ", "_")
    filename = filename.replace("ó", "o")
    filename = filename.replace("á", "a")
    filename = filename.replace("é", "e")
    filename = filename.replace("í", "i")
    filename = filename.replace("ú", "u")

    return f"recipes/{filename}.json"


def load_processes():

    path = "recipes/processes.json"

    if not os.path.exists(path):

        with open(path, "w", encoding="utf-8") as f:
            json.dump([], f)

    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_processes(processes):

    with open(
        "recipes/processes.json",
        "w",
        encoding="utf-8"
    ) as f:

        json.dump(
            processes,
            f,
            indent=4,
            ensure_ascii=False
        )


def tpm_database_path():

    os.makedirs(app.instance_path, exist_ok=True)
    path = os.path.join(app.instance_path, "tpm_reports.sqlite3")

    with closing(sqlite3.connect(path)) as connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS tpm_reports (
                report_number INTEGER PRIMARY KEY AUTOINCREMENT,
                machine TEXT NOT NULL,
                report_date TEXT NOT NULL,
                report_time TEXT NOT NULL,
                ok_count INTEGER NOT NULL,
                not_ok_count INTEGER NOT NULL,
                responses_json TEXT NOT NULL,
                maintenance_order TEXT NOT NULL DEFAULT ''
            )
        """)
        columns = {
            row[1]
            for row in connection.execute("PRAGMA table_info(tpm_reports)")
        }
        if "maintenance_order" not in columns:
            connection.execute("""
                ALTER TABLE tpm_reports
                ADD COLUMN maintenance_order TEXT NOT NULL DEFAULT ''
            """)
        connection.commit()

    return path


# ----------------------------------
# Home
# ----------------------------------

@app.route("/")
def home():
    return render_template("index.html")


# ----------------------------------
# Process Guide
# ----------------------------------

@app.route("/process")
def process():

    processes = load_processes()

    return render_template(
        "process.html",
        processes=processes
    )


# ----------------------------------
# Add Process
# ----------------------------------

@app.route("/add-process", methods=["POST"])
def add_process():

    name = request.form["process_name"].strip()

    if not name:
        return redirect("/process")

    processes = load_processes()

    if name not in processes:

        processes.append(name)

        save_processes(processes)

        filepath = recipe_file(name)

        recipe = {
            "process": name,
            "steps": []
        }

        with open(filepath, "w", encoding="utf-8") as f:

            json.dump(
                recipe,
                f,
                indent=4,
                ensure_ascii=False
            )

    return redirect("/process")


# ----------------------------------
# Delete Process
# ----------------------------------

@app.route("/delete-process", methods=["POST"])
def delete_process():

    process_name = request.form["process_name"]

    processes = load_processes()

    if process_name in processes:

        processes.remove(process_name)

        save_processes(processes)

        filepath = recipe_file(process_name)

        if os.path.exists(filepath):
            os.remove(filepath)

    return redirect("/process")


# ----------------------------------
# Recipe Editor
# ----------------------------------
@app.route("/recipe-editor", methods=["GET", "POST"])
def recipe_editor():

    process = request.args.get(
        "process",
        "Aceleración F585"
    )

    filepath = recipe_file(process)

    if os.path.exists(filepath):

        with open(filepath, "r", encoding="utf-8") as f:
            recipe = json.load(f)

    else:

        recipe = {
            "process": process,
            "steps": []
        }

    if request.method == "POST":

        titles = request.form.getlist("title")
        times = request.form.getlist("time")

        steps = []

        for i in range(len(titles)):

            # Keep existing media by default
            media_filename = ""

            if i < len(recipe["steps"]):
                media_filename = recipe["steps"][i].get(
                    "media",
                    ""
                )

            # Check if a new file was uploaded
            media_key = f"media{i}"

            if media_key in request.files:

                file = request.files[media_key]

                if file and file.filename:

                    media_filename = secure_filename(
                        file.filename
                    )

                    file.save(
                        os.path.join(
                            "static/media",
                            media_filename
                        )
                    )

            steps.append({
                "title": titles[i],
                "time": int(times[i]),
                "media": media_filename
            })

        recipe["steps"] = steps

        with open(filepath, "w", encoding="utf-8") as f:

            json.dump(
                recipe,
                f,
                indent=4,
                ensure_ascii=False
            )

        return redirect(
            f"/recipe-editor?process={process}"
        )

    return render_template(
        "recipe_editor.html",
        process=process,
        recipe=recipe
    )

# ----------------------------------
# Player
# ----------------------------------

@app.route("/player")
def player():

    process = request.args.get("process")

    filepath = recipe_file(process)

    if not os.path.exists(filepath):
        return redirect("/process")

    with open(filepath, "r", encoding="utf-8") as f:
        recipe = json.load(f)

    return render_template(
        "player.html",
        recipe=recipe
    )

# ----------------------------------
# Blink
# ----------------------------------
def set_blink_led(is_on):

    if GPIO is not None:
        state = GPIO.LOW if is_on else GPIO.HIGH
        GPIO.output(BLINK_PIN, state)
    else:
        state = "ON" if is_on else "OFF"
        print(f"GPIO {BLINK_PIN} {state} (simulado)")


def blink_led():

    for _ in range(1):

        set_blink_led(is_on=True)

        time.sleep(0.7)

        set_blink_led(is_on=False)

        time.sleep(0.2)
@app.route("/blink-step")
def blink_step():

    threading.Thread(
        target=blink_led
    ).start()

    return "OK"

@app.route("/process-complete")
def process_complete():

    set_blink_led(is_on=True)
    time.sleep(1)
    set_blink_led(is_on=False)

    return "OK"


# ----------------------------------
# TPM
# ----------------------------------

def tpm_checklist_path(mill):
    return os.path.join(
        app.root_path,
        "recipes",
        f"tpm_molino_{mill}.json"
    )


def default_tpm_items(mill):
    if mill != 2:
        return []

    return [
        {
            "description": "Paros de emergencia (12 paros en total) en buen estado, sin daños visibles",
            "details": "Presionar el paro de emergencia y verificar que la máquina no realice ningún movimiento.",
            "images": ["M2TPM_1.png"]
        },
        {
            "description": "Barra de emergencia en buen estado, sin daños visibles y libre de polvo",
            "details": "Presionar la barra de emergencia y verificar que la máquina no realice ningún movimiento. La máquina debe parar sin tener que aplicar una fuerza mayor a 40 lb.",
            "images": ["M2TPM_2.png"]
        },
        {
            "description": "Bases de cuchillas libres de daños visibles",
            "details": "",
            "images": ["M2TPM_3.png"]
        },
        {
            "description": "Balatas en buen estado, sin daños visibles o apertura anormal",
            "details": "",
            "images": ["M2TPM_4.png"]
        },
        {
            "description": "Motor en buen estado, sin daños, ruidos extraños y libre de polvo",
            "details": "",
            "images": ["M2TPM_5.png"]
        },
        {
            "description": "Llave hidroneumática en buen estado, sin daños visibles y libre de polvo",
            "details": "",
            "images": ["M2TPM_6.png"]
        },
        {
            "description": "Panel de control en buen estado, sin daños visibles y libre de polvo",
            "details": "",
            "images": ["M2TPM_7.png"]
        },
        {
            "description": "Medidor de temperatura de agua en buen estado, sin daños visibles y libre de polvo",
            "details": "Verificar que la luz se encuentre en color verde",
            "images": ["M2TPM_8.png", "M2TPM_8.1.png"]
        },
        {
            "description": "Rodillos en buen estado, sin daños visibles y libre de polvo",
            "details": "",
            "images": ["M2TPM_9.png"]
        },
        {
            "description": "Distancia de paro menor o igual a la distancia calculada",
            "details": "",
            "images": ["M2TPM_10.png"]
        },
        {
            "description": "Cilindros de apertura en buen estado, sin daños, ruidos extraños o fugas de aceite",
            "details": "",
            "images": ["M2TPM_11.png"]
        },
        {
            "description": "Niveles de aceite por arriba de la marca roja. Revisar que no existan fugas de aceite",
            "details": "",
            "images": ["M2TPM_12.png", "M2TPM_12.1.png"]
        },
        {
            "description": "Medición de la temperatura de mesa de enfriamiento de 15° a 25°",
            "details": "",
            "images": ["M2TPM_13.png"]
        }
    ]


def load_tpm_items(mill):
    path = tpm_checklist_path(mill)
    if not os.path.exists(path):
        items = default_tpm_items(mill)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8") as file:
            json.dump(items, file, indent=4, ensure_ascii=False)
        return items

    with open(path, "r", encoding="utf-8") as file:
        items = json.load(file)

    for item in items:
        images = item.get("images", [])
        if isinstance(images, str):
            images = [images]
        item["images"] = [
            image for image in images
            if isinstance(image, str) and os.path.basename(image) == image
        ]
        item.setdefault("details", "")

    return items


def render_tpm_checklist(mill, items, selected_statuses=None, form_error=None, status=200):
    return render_template(
        "tpm_checklist.html",
        mill=mill,
        items=items,
        selected_statuses=selected_statuses or {},
        form_error=form_error
    ), status


@app.route("/tpm")
def tpm():
    with closing(sqlite3.connect(tpm_database_path())) as connection:
        reports = connection.execute("""
            SELECT machine, report_date, report_time
            FROM tpm_reports
            ORDER BY report_number DESC
        """).fetchall()

    latest_reports = {}
    for machine, report_date, report_time in reports:
        latest_reports.setdefault(machine, (report_date, report_time))

    return render_template("tpm.html", latest_reports=latest_reports)


@app.route("/tpm-molino-<int:mill>", methods=["GET", "POST"])
def tpm_molino(mill):
    if mill not in (1, 2, 3):
        return "Molino no encontrado", 404

    items = load_tpm_items(mill)
    item_names = [f"item_{number}" for number in range(1, len(items) + 1)]

    if request.method == "POST":
        responses = {
            name: request.form.get(name, "")
            for name in item_names
        }
        allowed_responses = {"OK", "Not OK"}

        if not items:
            return render_tpm_checklist(
                mill,
                items,
                form_error="Este TPM todavía no tiene puntos configurados.",
                status=400
            )

        if any(response not in allowed_responses for response in responses.values()):
            return render_tpm_checklist(
                mill,
                items,
                selected_statuses=responses,
                form_error="Selecciona OK o Not OK para cada punto antes de enviar.",
                status=400
            )

        ok_count = sum(response == "OK" for response in responses.values())
        not_ok_count = sum(response == "Not OK" for response in responses.values())
        now = datetime.now()

        with closing(sqlite3.connect(tpm_database_path())) as connection:
            cursor = connection.execute("""
                INSERT INTO tpm_reports (
                    machine, report_date, report_time, ok_count, not_ok_count, responses_json
                ) VALUES (?, ?, ?, ?, ?, ?)
            """, (
                f"Molino {mill}",
                now.strftime("%d/%m/%Y"),
                now.strftime("%H:%M"),
                ok_count,
                not_ok_count,
                json.dumps(responses, ensure_ascii=False)
            ))
            report_number = cursor.lastrowid
            connection.commit()

        if not_ok_count:
            return render_template(
                "tpm_maintenance_order.html",
                report_number=report_number
            )

        return redirect("/tpm")

    return render_tpm_checklist(mill, items)


@app.route("/tpm/maintenance-order", methods=["POST"])
def save_tpm_maintenance_order():
    report_number = request.form.get("report_number", "")
    maintenance_order = request.form.get("maintenance_order", "").strip()

    if not report_number.isdigit() or len(maintenance_order) > 120:
        return redirect("/tpm")

    with closing(sqlite3.connect(tpm_database_path())) as connection:
        connection.execute("""
            UPDATE tpm_reports
            SET maintenance_order = ?
            WHERE report_number = ? AND not_ok_count > 0
        """, (maintenance_order, int(report_number)))
        connection.commit()

    return redirect("/tpm")


@app.route("/tpm/reports")
def tpm_reports():
    with closing(sqlite3.connect(tpm_database_path())) as connection:
        reports = connection.execute("""
            SELECT report_number, machine, report_date, report_time, ok_count, not_ok_count
            FROM tpm_reports
            ORDER BY report_number DESC
        """).fetchall()

    return render_template("tpm_reports.html", reports=reports)


@app.route("/tpm/download-report")
def download_tpm_report():
    with closing(sqlite3.connect(tpm_database_path())) as connection:
        reports = connection.execute("""
            SELECT report_number, machine, report_date, report_time,
                   ok_count, not_ok_count, responses_json, maintenance_order
            FROM tpm_reports
            ORDER BY report_number
        """).fetchall()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "TPM Reports"

    responses_by_report = [json.loads(report[6]) for report in reports]
    response_count = max(
        (len(responses) for responses in responses_by_report),
        default=13
    )
    response_count = max(response_count, 13)
    headers = [
        "N. reporte", "TPM", "Fecha", "Hora", "Resp. OK", "Resp. NOK",
        *(f"R{number}" for number in range(1, response_count + 1)),
        "Orden de mantenimiento"
    ]
    worksheet.append(headers)

    header_fill = PatternFill("solid", fgColor="B4C7E7")
    ok_fill = PatternFill("solid", fgColor="E2F0D9")
    not_ok_fill = PatternFill("solid", fgColor="F4CCCC")
    thin_gray = Side(style="thin", color="D9E1F2")
    cell_border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left")
        cell.border = cell_border

    for report, responses in zip(reports, responses_by_report):
        row = [
            report[0],
            report[1],
            datetime.strptime(report[2], "%d/%m/%Y").date(),
            datetime.strptime(report[3], "%H:%M").time(),
            report[4],
            report[5],
            *[
                (
                    "ok" if responses.get(f"item_{number}") == "OK"
                    else "nok" if responses.get(f"item_{number}") == "Not OK"
                    else ""
                )
                for number in range(1, response_count + 1)
            ],
            report[7]
        ]
        worksheet.append(row)
        current_row = worksheet.max_row
        row_fill = not_ok_fill if report[5] > 0 else ok_fill

        for cell in worksheet[current_row]:
            cell.fill = row_fill
            cell.border = cell_border
            cell.alignment = Alignment(vertical="center")

        worksheet.cell(current_row, 1).number_format = "0000"
        worksheet.cell(current_row, 3).number_format = "DD/MM/YYYY"
        worksheet.cell(current_row, 4).number_format = "HH:MM"
        for column in range(5, 7 + response_count):
            worksheet.cell(current_row, column).alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    column_widths = [14, 16, 14, 10, 12, 12] + [8] * response_count + [30]
    for column_number, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[
            worksheet.cell(row=1, column=column_number).column_letter
        ].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="tpm_reports.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ----------------------------------
# Settings
# ----------------------------------

@app.route("/settings")
def settings():
    return render_template("settings.html")


@app.route("/settings/tpm/<int:mill>", methods=["GET", "POST"])
def tpm_checklist_editor(mill):
    if mill not in (1, 2, 3):
        return "Molino no encontrado", 404

    items = load_tpm_items(mill)
    error = None

    if request.method == "POST":
        descriptions = request.form.getlist("description")
        details = request.form.getlist("details")
        updated_items = []
        uploads = []

        for index, description in enumerate(descriptions):
            description = description.strip()
            detail = details[index].strip() if index < len(details) else ""
            if not description:
                if detail or request.files.getlist(f"images{index}"):
                    error = "Cada punto con instrucciones o imagen necesita una descripción."
                    break
                continue

            existing_images = request.form.getlist(f"existing_image{index}")
            existing_images = [
                image for image in existing_images
                if os.path.basename(image) == image
            ]
            remove_images = set(request.form.getlist(f"remove_image{index}"))
            kept_images = [
                image for image in existing_images
                if os.path.basename(image) == image and image not in remove_images
            ]
            item_images = list(kept_images)

            for uploaded_file in request.files.getlist(f"images{index}"):
                if not uploaded_file.filename:
                    continue
                extension = os.path.splitext(
                    secure_filename(uploaded_file.filename)
                )[1].lower()
                if extension not in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
                    error = "Solo se permiten imágenes PNG, JPG, JPEG, GIF o WEBP."
                    break
                filename = f"m{mill}_{uuid.uuid4().hex}{extension}"
                item_images.append(filename)
                uploads.append((uploaded_file, filename))

            if error:
                break

            updated_items.append({
                "description": description,
                "details": detail,
                "images": item_images
            })

        if error is None:
            media_directory = os.path.join(
                app.root_path,
                "static",
                "images",
                "TPM"
            )
            os.makedirs(media_directory, exist_ok=True)
            for uploaded_file, filename in uploads:
                uploaded_file.save(os.path.join(media_directory, filename))

            with open(tpm_checklist_path(mill), "w", encoding="utf-8") as file:
                json.dump(updated_items, file, indent=4, ensure_ascii=False)

            return redirect(f"/settings/tpm/{mill}")

    return render_template(
        "tpm_checklist_editor.html",
        mill=mill,
        items=items,
        error=error
    ), 400 if error else 200


# ----------------------------------
# Run
# ----------------------------------

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=5004,
        debug=True
    )
